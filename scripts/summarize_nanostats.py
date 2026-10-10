#!/usr/bin/env python3
"""Create TSV and XLSX summaries from NanoPlot NanoStats.txt files."""

from __future__ import annotations

import argparse
import csv
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any


METRICS = (
    "Mean read length",
    "Mean read quality",
    "Median read length",
    "Median read quality",
    "Number of reads",
    "Read length N50",
)
HEADERS = ("Sample", *METRICS)
INTEGER_METRICS = {"Number of reads", "Read length N50"}


def log(message: str) -> None:
    """Print a timestamped message for the pipeline log."""
    stamp = datetime.now().astimezone().strftime("%A %Y-%m-%d %H:%M:%S %Z")
    print(f"[{stamp}] {message}", flush=True)


def natural_sort_key(value: str) -> list[object]:
    """Sort sample2 before sample10."""
    return [
        int(part) if part.isdigit() else part.casefold()
        for part in re.split(r"(\d+)", value)
    ]


def numeric_value(raw_value: str, metric: str) -> int | float | str:
    """Convert a NanoStats value to an Excel-friendly number when possible."""
    cleaned = raw_value.strip().replace(",", "")
    match = re.search(
        r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?",
        cleaned,
    )

    if not match:
        return raw_value.strip()

    number = float(match.group())
    if metric in INTEGER_METRICS and number.is_integer():
        return int(number)
    return number


def read_nanostats(stats_file: Path) -> dict[str, Any]:
    """Extract the requested metrics from one NanoStats.txt file."""
    content = stats_file.read_text(encoding="utf-8", errors="replace")
    results: dict[str, Any] = {}

    for metric in METRICS:
        pattern = re.compile(
            rf"^\s*{re.escape(metric)}\s*:\s*(.*?)\s*$",
            flags=re.IGNORECASE | re.MULTILINE,
        )
        match = pattern.search(content)
        results[metric] = numeric_value(match.group(1), metric) if match else None

    return results


def collect_rows(input_dir: Path) -> list[dict[str, Any]]:
    """Collect one summary row per NanoPlot sample directory."""
    stats_files = sorted(
        input_dir.rglob("NanoStats.txt"),
        key=lambda path: natural_sort_key(str(path)),
    )

    if not stats_files:
        raise FileNotFoundError(f"No NanoStats.txt files were found in {input_dir}.")

    rows: list[dict[str, Any]] = []
    sample_sources: dict[str, Path] = {}

    for stats_file in stats_files:
        sample = stats_file.parent.name
        normalized_sample = sample.casefold()

        if normalized_sample in sample_sources:
            raise ValueError(
                f"Multiple NanoStats.txt files were found for {sample}: "
                f"{sample_sources[normalized_sample]} and {stats_file}."
            )

        sample_sources[normalized_sample] = stats_file
        metrics = read_nanostats(stats_file)
        row: dict[str, Any] = {"Sample": sample, **metrics}
        rows.append(row)

        missing = [metric for metric in METRICS if metrics[metric] is None]
        if missing:
            log(
                f"WARNING: {sample} is missing these metrics: "
                + ", ".join(missing)
            )

        log(f"Read statistics for {sample}: {stats_file}")

    return sorted(rows, key=lambda row: natural_sort_key(str(row["Sample"])))


def write_tsv(rows: list[dict[str, Any]], output_file: Path) -> None:
    """Write the summary as a tab-separated table."""
    with output_file.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS, delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)


def write_excel(rows: list[dict[str, Any]], output_file: Path) -> None:
    """Write a formatted Excel workbook."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as error:
        raise RuntimeError(
            "openpyxl is required. Activate environment.yml or install "
            "requirements.txt."
        ) from error

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "NanoPlot QC"
    sheet.append(list(HEADERS))

    for row in rows:
        sheet.append([row.get(header) for header in HEADERS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill

    sheet.freeze_panes = "A2"

    for row_number in range(2, sheet.max_row + 1):
        for column_number in range(2, 6):
            sheet.cell(row=row_number, column=column_number).number_format = "0.00"
        for column_number in range(6, 8):
            sheet.cell(row=row_number, column=column_number).number_format = "#,##0"

    for column_number, header in enumerate(HEADERS, start=1):
        values = [
            str(sheet.cell(row=row, column=column_number).value or "")
            for row in range(1, sheet.max_row + 1)
        ]
        width = min(max(len(value) for value in values) + 2, 35)
        sheet.column_dimensions[get_column_letter(column_number)].width = width

    end_column = get_column_letter(len(HEADERS))
    table = Table(
        displayName="NanoPlotSummary",
        ref=f"A1:{end_column}{sheet.max_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    workbook.save(output_file)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Build TSV and XLSX tables from NanoPlot NanoStats.txt files."
    )
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--prefix", default="nanoplot_summary")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    input_dir = args.input_dir.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()

    if not input_dir.is_dir():
        log(f"ERROR: input directory does not exist: {input_dir}")
        return 2

    output_dir.mkdir(parents=True, exist_ok=True)
    tsv_output = output_dir / f"{args.prefix}.tsv"
    excel_output = output_dir / f"{args.prefix}.xlsx"

    try:
        rows = collect_rows(input_dir)
        write_tsv(rows, tsv_output)
        write_excel(rows, excel_output)
    except (FileNotFoundError, OSError, RuntimeError, ValueError) as error:
        log(f"ERROR: {error}")
        return 1

    log(f"TSV summary created: {tsv_output}")
    log(f"Excel summary created: {excel_output}")
    log(f"Summary contains {len(rows)} sample(s).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
