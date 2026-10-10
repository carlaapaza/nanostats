import tempfile
import unittest
from pathlib import Path

from scripts.run_nanoplot import group_reads, strip_read_suffix
from scripts.summarize_nanostats import collect_rows, read_nanostats, write_excel


NANOSTATS_TEXT = """General summary:
Mean read length:             1,298.6
Mean read quality:            12.9
Median read length:           464.0
Median read quality:          13.5
Number of reads:              380,928.0
Read length N50:              2,731.0
"""


class PipelineTests(unittest.TestCase):
    def test_supported_suffixes_are_removed(self):
        self.assertEqual(strip_read_suffix("barcode01.fastq.gz"), "barcode01")
        self.assertEqual(strip_read_suffix("sample.fq.bz2"), "sample")
        self.assertEqual(strip_read_suffix("sample.FASTQ"), "sample")

    def test_barcode_directory_groups_multiple_chunks(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            input_dir = Path(temporary_directory)
            barcode_dir = input_dir / "barcode02"
            barcode_dir.mkdir()
            first = barcode_dir / "chunk01.fastq.gz"
            second = barcode_dir / "chunk02.fastq.gz"
            first.touch()
            second.touch()

            groups = group_reads([first.resolve(), second.resolve()], input_dir)

            self.assertEqual(list(groups), ["barcode02"])
            self.assertEqual(len(groups["barcode02"]), 2)

    def test_nanostats_metrics_are_numeric(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            stats_file = Path(temporary_directory) / "NanoStats.txt"
            stats_file.write_text(NANOSTATS_TEXT, encoding="utf-8")

            metrics = read_nanostats(stats_file)

            self.assertEqual(metrics["Mean read length"], 1298.6)
            self.assertEqual(metrics["Number of reads"], 380928)
            self.assertEqual(metrics["Read length N50"], 2731)

    def test_summary_rows_use_natural_sample_order(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            input_dir = Path(temporary_directory)

            for barcode in ("barcode10", "barcode2"):
                sample_dir = input_dir / barcode
                sample_dir.mkdir()
                (sample_dir / "NanoStats.txt").write_text(
                    NANOSTATS_TEXT,
                    encoding="utf-8",
                )

            rows = collect_rows(input_dir)

            self.assertEqual(
                [row["Sample"] for row in rows],
                ["barcode2", "barcode10"],
            )

    def test_excel_summary_is_created(self):
        from openpyxl import load_workbook

        row = {"Sample": "barcode01"}
        row.update(read_nanostats_from_text(NANOSTATS_TEXT))

        with tempfile.TemporaryDirectory() as temporary_directory:
            output_file = Path(temporary_directory) / "summary.xlsx"
            write_excel([row], output_file)

            workbook = load_workbook(output_file, data_only=True)
            sheet = workbook["NanoPlot QC"]

            self.assertEqual(sheet.max_row, 2)
            self.assertEqual(sheet["A2"].value, "barcode01")
            self.assertEqual(sheet["F2"].value, 380928)


def read_nanostats_from_text(content):
    with tempfile.TemporaryDirectory() as temporary_directory:
        stats_file = Path(temporary_directory) / "NanoStats.txt"
        stats_file.write_text(content, encoding="utf-8")
        return read_nanostats(stats_file)


if __name__ == "__main__":
    unittest.main()
